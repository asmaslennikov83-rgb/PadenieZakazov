import asyncio
from contextlib import closing
from datetime import date,datetime,timedelta,timezone
from io import BytesIO
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import AsyncMock,patch
import zipfile
import openpyxl
from history_store import Store,key,now,dumps
from business_report import build,classify,change
from analytical_export import build_package
from collector import Collector,ReadOnlyHTTP,SourceError,MSK
from manual_import import import_rows
from maintenance import restore

DAY=date(2026,10,19)


def seed(store,c='spb_tt',n='123',base=2,fact=1,days=90,created=None):
    start=DAY-timedelta(days=days-1)
    with store.connect() as db:store.product(db,c,n,{'supplier_article':'SKU '+n,'title':'Товар','created_at':created},now())
    records=[]
    for offset in range(days):
        d=start+timedelta(days=offset)
        qty=fact if d==DAY else base
        for i in range(qty):
            records.append({'event_key':key(n,str(d),i),'nm_id':n,'day':str(d),'units':1,'order_price':100,'buyer_price':90,'is_cancel':False,'model':'FBO' if i%2 else 'FBS','barcode':str(i%2),'size':'M','event_type':'order','product':{}})
    store.events('fixture',c,'orders',records)
    store.cover('fixture',c,'orders',start,DAY)
    store.source('fixture',c,'orders',status='OK',reason='Fixture, not real WB',requested_at=now(),completed_at=now(),pages=2,records=len(records))


class ModernTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.path=Path(self.temp.name)/'bot.db';self.store=Store(self.path);self.store.migrate()
    def tearDown(self):self.temp.cleanup()
    def report(self,codes=None):return build(self.store,DAY,codes or ['spb_tt'],save=False)

    def test_migration_preserves_existing_settings_and_backup_restores(self):
        p=Path(self.temp.name)/'old.db'
        with closing(sqlite3.connect(p)) as db:
            db.execute('CREATE TABLE user_settings(user_id INTEGER PRIMARY KEY,threshold INTEGER,report_time TEXT)')
            db.execute("INSERT INTO user_settings VALUES(42,25,'10:00')")
            db.commit()
        store=Store(p);backup=store.migrate();self.assertTrue(backup.exists())
        self.assertEqual(store.rows('SELECT threshold FROM user_settings')[0]['threshold'],25)
        self.assertIsNone(store.migrate())
        restore(store,backup)
        self.assertEqual(store.rows('SELECT report_time FROM user_settings')[0]['report_time'],'10:00')
        self.assertFalse(store.rows("SELECT name FROM sqlite_master WHERE name='v3_products'"))

    def test_zero_orders_with_complete_source_is_real_zero(self):
        seed(self.store,fact=0)
        r=self.report();p=r['all_products'][0]
        self.assertEqual(p['ordered_units'],0);self.assertEqual(p['orders_status'],'OK')
        self.assertEqual(r['summary'][0]['change_rate'],-1)

    def test_today_is_not_compared_to_complete_day(self):
        with self.assertRaises(ValueError):build(self.store,datetime.now(MSK).date(),['spb_tt'])

    def test_missing_source_never_becomes_full_zero(self):
        with self.store.connect() as db:self.store.product(db,'spb_tt','123',{},now())
        r=self.report();self.assertIsNone(r['all_products'][0]['ordered_units']);self.assertEqual(r['status'],'PARTIAL')

    def test_failure_one_cabinet_preserves_other_and_combined_partial(self):
        seed(self.store);self.store.source('bad','spb_tt_group','orders',status='NO_ACCESS',reason='HTTP 403',records=0)
        r=self.report(['spb_tt','spb_tt_group'])
        self.assertEqual(r['summary'][0]['ordered_units'],1);self.assertIsNone(r['summary'][1]['ordered_units']);self.assertEqual(r['status'],'PARTIAL')

    def test_zero_base_growth_without_infinite_percentage(self):
        seed(self.store,base=0,fact=3);p=self.report()['all_products'][0]
        self.assertEqual(p['signal'],'ПОЯВИЛИСЬ ЗАКАЗЫ');self.assertIsNone(p['change_rate'])
        self.assertEqual(classify(0,0),'НЕТ ЗАКАЗОВ')

    def test_rare_demand_has_separate_signal(self):
        seed(self.store,base=0,fact=0)
        for i in range(7):
            self.store.events('extra','spb_tt','orders',[{'event_key':f'rare{i}','nm_id':'123','day':str(DAY-timedelta(days=i+1)),'units':1,'order_price':100,'buyer_price':90,'is_cancel':False,'model':'FBO','barcode':'X','size':'M','event_type':'order'}])
        p=self.report()['all_products'][0];self.assertEqual(p['base_14_avg'],.5);self.assertEqual(p['signal'],'РЕДКИЙ СПРОС')

    def test_new_card_incomplete_base_and_union_reconciles(self):
        seed(self.store,base=0,fact=4,created='2026-10-18T10:00:00+03:00')
        r=self.report();self.assertIsNone(r['all_products'][0]['base_14_avg'])
        self.assertEqual(r['summary'][0]['ordered_units'],4);self.assertAlmostEqual(r['summary'][0]['reconciliation_delta'],0)

    def test_gross_drop_and_gain_compensation(self):
        seed(self.store,n='123',base=10,fact=2);seed(self.store,n='124',base=2,fact=10)
        s=self.report()['summary'][0]
        self.assertEqual(s['gross_drop'],8);self.assertEqual(s['gross_gain'],8);self.assertEqual(s['delta_units'],0);self.assertEqual(s['change_rate'],0)

    def test_variant_rows_do_not_duplicate_card_units(self):
        seed(self.store,base=4,fact=3);r=self.report()
        vd=[x for x in r['tables']['VariantDaily'] if x['day']==DAY]
        self.assertEqual(sum(x['ordered_units'] for x in vd),r['all_products'][0]['ordered_units'])
        self.assertTrue(all('views' not in x for x in vd))

    def test_identical_seller_article_across_cabinets_never_auto_matches(self):
        seed(self.store,c='spb_tt');seed(self.store,c='spb_tt_group')
        r=self.report(['spb_tt','spb_tt_group']);self.assertEqual(len(r['all_products']),2);self.assertEqual(r['comparisons'],[])

    def test_confirmed_pack_mapping_normalizes_units(self):
        seed(self.store,c='spb_tt',fact=2);seed(self.store,c='spb_tt_group',fact=1)
        rows=[dict(canonical_product_id='P',cabinet_id=c,nm_id='123',variant_id='',units_per_item=u,comparison_unit='piece',mapping_source='Owner confirmed',mapping_status='CONFIRMED') for c,u in [('spb_tt','3'),('spb_tt_group','1')]]
        import_rows(self.store,'ProductMapping',rows)
        r=self.report(['spb_tt','spb_tt_group']);self.assertEqual(r['comparisons'][0]['fact_units'],7)

    def test_future_stock_not_used_for_historical_report(self):
        seed(self.store)
        row={'observation_key':'x','nm_id':'123','quantity':100,'variant_id':'v','warehouse_id':'w','model':'FBO'}
        self.store.observations('spb_tt','stocks_wb',[row],'2026-10-20T10:00:00+03:00')
        self.assertIsNone(self.report()['all_products'][0]['fbo_snapshot_quantity'])

    def test_utc_snapshot_after_moscow_midnight_is_not_previous_day(self):
        seed(self.store)
        self.store.observations('spb_tt','stocks_wb',[{'observation_key':'x','nm_id':'123','quantity':100}], '2026-10-19T23:30:00+00:00')
        self.assertIsNone(self.report()['all_products'][0]['fbo_snapshot_quantity'])

    def test_instance_lock_blocks_parallel_process_state(self):
        from runtime_lock import instance_lock
        path=Path(self.temp.name)/'instance.lock'
        with instance_lock(path):
            with self.assertRaises(RuntimeError):
                with instance_lock(path):pass
        with instance_lock(path):pass

    def test_sparse_zero_snapshots_do_not_create_outage_hours(self):
        seed(self.store)
        for stamp in ('2026-10-18T10:00:00+03:00','2026-10-19T10:00:00+03:00'):
            self.store.observations('spb_tt','stocks_wb',[{'observation_key':'x','nm_id':'123','quantity':0}],stamp)
        p=self.report()['all_products'][0]
        self.assertNotIn('outage_hours',p);self.assertIsNone(p['stock_quantity'])

    def test_historical_csv_supplies_dated_stock_without_false_warehouse_id(self):
        from historical_stock_adapter import parse_details
        raw=BytesIO()
        with zipfile.ZipFile(raw,'w') as z:z.writestr('s.csv','NmID;ChrtID;OfficeName;18.10.2026;19.10.2026\n123;5;Склад WB;3;0\n')
        records=list(parse_details(raw.getvalue(),DAY-timedelta(days=1),DAY,'FBO','fixture'))
        self.assertEqual(len(records),2);self.assertEqual(records[1]['quantity'],0);self.assertIsNone(records[1]['warehouse_id'])
        seed(self.store)
        self.store.observations('spb_tt','stock_history_wb',records,None)
        self.store.source('h','spb_tt','stock_history_wb',status='OK',reason='Fixture')
        p=self.report()['all_products'][0];self.assertEqual(p['fbo_snapshot_quantity'],0);self.assertEqual(p['fbo_status'],'OK')
        self.assertIsNone(p['stock_quantity'])

    def test_overlapping_scenarios_are_not_automatically_added(self):
        seed(self.store)
        with self.store.connect() as db:
            for i in range(1,15):
                d=DAY-timedelta(days=i);old=i>7
                obj={'views':100,'carts':20,'funnel_orders':10 if old else 5,'wishlist':0,'status':'OK'}
                db.execute('INSERT INTO v3_funnel VALUES(?,?,?,?,?,?)',('spb_tt','123',str(d),dumps(obj),'f',now()))
        r=self.report();ev=[e for e in r['tables']['Evidence'] if e['rule_id']=='CONVERSION_RESTORE']
        self.assertTrue(ev);self.assertIsNone(ev[0]['scenario_units']);self.assertFalse(ev[0]['scenario_additive'])
        import_rows(self.store,'AvailabilityPeriods',[{'cabinet_id':'spb_tt','nm_id':'123','period_start':'2026-10-05','period_end':'2026-10-18','scope':'Card all variants compared period','availability_status':'CONFIRMED_AVAILABLE','source':'Owner verification','status':'OK'}])
        ev=[e for e in self.report()['tables']['Evidence'] if e['rule_id']=='CONVERSION_RESTORE'][0]
        self.assertEqual(ev['scenario_units'],35);self.assertFalse(ev['scenario_additive'])

    def test_no_margin_no_claim_of_profitability(self):
        seed(self.store,base=2,fact=8)
        ev=[e for e in self.report()['tables']['Evidence'] if e['rule_id']=='OBSERVED_GROWTH'][0]
        self.assertIn('экономику',ev['next_action']);self.assertIsNone(ev['scenario_units'])

    def test_ads_exceed_total_flag_attribution_without_negative_organic(self):
        seed(self.store)
        import_rows(self.store,'AdsDaily',[{'cabinet_id':'spb_tt','campaign_id':'1','nm_id':'123','day':str(DAY),'spend':'100','impressions':'1000','clicks':'10','attributed_orders':'99','attributed_amount':'200','attribution_window':'manual 7d click','source':'Owner export','status':'OK'}])
        ad=self.report()['tables']['AdsDaily'][0]
        self.assertEqual(ad['attribution_check'],'CHECK_WINDOW_AND_SCOPE');self.assertNotIn('organic_orders',ad);self.assertEqual(ad['drr'],.5)

    def test_price_and_decline_only_hypothesis(self):
        seed(self.store,base=2,fact=0)
        with self.store.connect() as db:
            for row in db.execute("SELECT event_key,payload FROM v3_events WHERE day>=?",(str(DAY-timedelta(days=7)),)).fetchall():
                obj=json.loads(row['payload']);obj['order_price']=200;db.execute('UPDATE v3_events SET payload=? WHERE event_key=?',(dumps(obj),row['event_key']))
        ev=[x for x in self.report()['tables']['Evidence'] if x['rule_id']=='PRICE_COINCIDENCE']
        self.assertTrue(ev);self.assertEqual(ev[0]['explanation_status'],'гипотеза')

    def test_reimport_and_event_revision_are_idempotent(self):
        seed(self.store)
        before=self.report()['summary'][0]['ordered_units'];seed(self.store)
        self.assertEqual(self.report()['summary'][0]['ordered_units'],before)

    def test_report_revisions_immutable_and_same_facts_reuse_revision(self):
        seed(self.store)
        first=build(self.store,DAY,['spb_tt']);second=build(self.store,DAY,['spb_tt'])
        self.assertEqual(first['revision'],second['revision'])
        files=list((self.path.parent/'report_snapshots').glob('*.json.gz'));self.assertEqual(len(files),1)
        self.store.source('bad','spb_tt','orders',status='ERROR',reason='Network')
        third=build(self.store,DAY,['spb_tt']);self.assertEqual(third['revision'],first['revision']+1)

    def test_delivery_claims_survive_restart_and_unknown_not_resent(self):
        self.assertTrue(self.store.delivery('daily:42'));self.store.delivery('daily:42','UNKNOWN')
        self.assertFalse(Store(self.path).delivery('daily:42'))

    def test_action_history_and_incomplete_result(self):
        seed(self.store);r=build(self.store,DAY,['spb_tt']);aid=r['tables']['Actions'][0]['action_id']
        self.store.action(aid,'DONE',owner='Owner')
        r=self.report();a=r['tables']['Actions'][0];self.assertEqual(a['status'],'DONE');self.assertEqual(a['result_status'],'MISSING')
        self.assertEqual(len(self.store.rows('SELECT * FROM v3_action_history')),1)

    def test_excel_cached_values_percent_types_and_schema(self):
        seed(self.store);r=self.report();name,blob=build_package(r)
        wb=openpyxl.load_workbook(BytesIO(blob),data_only=True)
        self.assertIn('Manifest',wb.sheetnames);self.assertIn('Dictionary',wb.sheetnames);self.assertIn('ProductDaily',wb.sheetnames)
        sheet=wb['Все товары'];self.assertEqual(sheet['B2'].data_type,'s');self.assertEqual(sheet['G2'].value,-.5);self.assertEqual(sheet['G2'].number_format,'0.0%')
        self.assertEqual(len(wb['Обзор']._charts),1);self.assertTrue(sheet.auto_filter.ref)
        self.assertEqual(wb['Все товары']['D2'].value,1)

    def test_full_zip_fallback_contains_all_rows_and_manifest(self):
        seed(self.store);r=self.report();name,blob=build_package(r,max_cells=1)
        self.assertTrue(name.endswith('.zip'))
        with zipfile.ZipFile(BytesIO(blob)) as z:
            self.assertIsNone(z.testzip());m=json.loads(z.read('manifest.json'))
            self.assertEqual(m['tables']['ProductDaily']['rows'],90)
            self.assertEqual(len(z.read('ProductDaily.csv').decode('utf-8-sig').splitlines()),91)

    def test_formula_injection_in_vendor_code_is_neutralized(self):
        seed(self.store)
        with self.store.connect() as db:self.store.product(db,'spb_tt','123',{'supplier_article':'=1+1'},now())
        _,data=build_package(self.report());w=openpyxl.load_workbook(BytesIO(data),data_only=False)
        self.assertEqual(w['Все товары']['C2'].data_type,'s')


class AsyncModernTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.store=Store(Path(self.temp.name)/'bot.db');self.store.migrate()
    def tearDown(self):self.temp.cleanup()
    async def test_statistics_paginates_to_empty_keeps_exact_fractional_cursor(self):
        today=datetime.now(MSK).date()
        obj={'srid':'private-order-token','nmId':123,'date':str(today-timedelta(days=1))+'T12:00:00','lastChangeDate':str(today)+'T10:00:00.1234567','priceWithDisc':100,'isCancel':False,'warehouseType':'Склад WB'}
        http=type('HTTP',(),{})();http.request=AsyncMock(side_effect=[[obj,obj],[]]);c=Collector(self.store,http)
        result=await c.sync_source('r','spb_tt','orders',lambda m:c.ledger('r','spb_tt','orders',m))
        self.assertEqual(result['status'],'OK');self.assertEqual(http.request.call_args_list[1].kwargs['params']['dateFrom'],obj['lastChangeDate'])
        events=self.store.rows('SELECT * FROM v3_events');self.assertEqual(len(events),1);self.assertNotIn('private-order-token',events[0]['payload'])

    async def test_page_interruption_retains_partial_and_no_fake_coverage(self):
        today=datetime.now(MSK).date()
        obj={'srid':'1','nmId':123,'date':str(today-timedelta(days=1))+'T12:00:00','lastChangeDate':str(today)+'T10:00:00','priceWithDisc':100,'isCancel':False}
        http=type('HTTP',(),{})();http.request=AsyncMock(side_effect=[[obj],SourceError('ERROR','HTTP 429')]);c=Collector(self.store,http)
        result=await c.sync_source('r','spb_tt','orders',lambda m:c.ledger('r','spb_tt','orders',m))
        self.assertEqual(result['status'],'PARTIAL');self.assertEqual(self.store.rows('SELECT * FROM v3_coverage'),[])
        self.assertTrue(self.store.sync_state('spb_tt','orders')['pending_cursor'])

    async def test_empty_stock_response_has_no_synthetic_zero_rows(self):
        http=type('HTTP',(),{})();http.request=AsyncMock(return_value={'data':{'items':[]}});c=Collector(self.store,http)
        result=await c.sync_source('r','spb_tt','stocks_wb',lambda m:c.stocks('r','spb_tt','wb',m))
        self.assertEqual(result['status'],'OK');self.assertEqual(self.store.rows('SELECT * FROM v3_observations'),[])

    async def test_read_only_guard_blocks_business_mutation(self):
        with self.assertRaises(ValueError):await ReadOnlyHTTP().request('spb_tt','bad','POST','https://advert-api.wildberries.ru/adv/v0/start','token')

    async def test_unauthorized_telegram_handler_does_not_get_report(self):
        import modern_ui
        event=type('Event',(),{'from_user':type('User',(),{'id':987654321})(),'answer':AsyncMock()})()
        with patch.object(modern_ui,'ALLOWED_TELEGRAM_IDS',set()):self.assertFalse(await modern_ui.allowed(event))
        event.answer.assert_awaited_once()

    async def test_ads_allocates_only_leaf_nm_platform_not_campaign_total(self):
        day=str(datetime.now(MSK).date()-timedelta(days=1))
        listing={'adverts':[{'status':9,'type':9,'advert_list':[{'advertId':1,'changeTime':now()}]}]}
        stats=[{'advertId':1,'currency':'RUB','days':[{'date':day+'T00:00:00Z','sum':30,'apps':[{'appType':32,'nms':[{'nmId':123,'sum':10,'views':100,'clicks':5,'orders':1,'shks':1,'sum_price':100},{'nmId':124,'sum':20,'views':200,'clicks':10,'orders':2,'shks':2,'sum_price':200}]}]}]}]
        http=type('HTTP',(),{})();http.request=AsyncMock(side_effect=[listing,stats]);c=Collector(self.store,http)
        self.store.sync_state('spb_tt','ads',{'history_complete':True})
        result=await c.sync_source('r','spb_tt','ads',lambda m:c.ads('r','spb_tt',m))
        self.assertEqual(result['status'],'OK')
        rows=[json.loads(r['payload']) for r in self.store.rows("SELECT payload FROM v3_observations WHERE source='ads'")]
        self.assertEqual(sum(r['spend'] for r in rows),30);self.assertEqual(len(rows),2)


if __name__=='__main__':unittest.main()
