"""One self-contained XLSX or ZIP; typed values, stable machine tables, no raw buyers."""
from __future__ import annotations
from collections import defaultdict
from datetime import date,datetime
from io import BytesIO,StringIO
import csv
import hashlib
import json
import os
import zipfile
from openpyxl import Workbook
from openpyxl.chart import LineChart,Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo
from metric_registry import TABLES


LABELS={
'cabinet_id':'Кабинет','nm_id':'Артикул WB','supplier_article':'Артикул продавца','title':'Название',
'ordered_units':'Заказано, шт.','known_ordered_units':'Известный подытог, шт.','orders_status':'Качество заказов',
'base_14_avg':'База 14 дней, шт./день','delta_units':'Отклонение, шт.','change_rate':'Изменение, %','signal':'Наблюдение',
'order_amount':'Сумма заказов, ₽','gross_drop':'Валовой спад, шт.','gross_gain':'Валовой прирост, шт.',
'available_days':'Доступно дней','expected_days':'Ожидается дней','status':'Статус','reason':'Причина/ограничение',
'fact':'Наблюдаемый факт','next_action':'Предлагаемая проверка','explanation_status':'Статус объяснения',
'scenario_units':'Сценарий, дополнительные шт.','scenario_limits':'Ограничения сценария','scenario_kind':'Вид сценария',
'scenario_horizon_days':'Горизонт, дней','rule_id':'Правило','evidence_id':'ID доказательства','action_id':'ID действия',
'source':'Источник','records':'Записей','completed_at':'Обновлено (ISO с поясом)','requested_at':'Запрошено (ISO с поясом)',
'report_date':'Отчётная дата','fbo_snapshot_quantity':'FBO, последний срез','fbs_snapshot_quantity':'FBS, последний срез',
'fbo_status':'Качество FBO','fbs_status':'Качество FBS','stock_coverage_days_estimate':'Оценка запаса, дней',
'weekday_avg':'Среднее сопоставимых дней','weekday_available':'Сопоставимых дней','weekday_expected':'Целевых дней',
'orders_last_7_sum':'Заказано за 7 дней','orders_previous_7_sum':'За предыдущие 7 дней','week_change_rate':'Недельное изменение, %',
'share_gross_drop':'Доля в валовом спаде','share_gross_gain':'Доля в приросте','first_observed_at':'Первое наблюдение',
'life_status':'Жизненный статус','owner':'Ответственный','due_date':'Срок проверки','metric_success':'Метрика успеха',
'canonical_product_id':'Подтверждённая группа','comparison_unit':'Единица сравнения','fact_units':'Заказано по паре','base_units':'База пары','cabinet_a':'Кабинет А','cabinet_b':'Кабинет Б','nm_a':'Артикул WB А','nm_b':'Артикул WB Б','limit':'Ограничение',
}
PERCENT_FIELDS={'change_rate','week_change_rate','weekday_change_rate','share_gross_drop','share_gross_gain','cart_rate','order_rate','seller_discount','ctr','drr'}
DEFINITIONS={
'base_14_avg':'Сумма заказанных единиц за D−14…D−1 / 14. Для карточки нужна полная применимая база; до подтверждённого создания карточки NOT_APPLICABLE.',
'delta_units':'Факт завершённого дня D минус base_14_avg; отрицательное значение означает снижение.',
'change_rate':'Факт / база − 1 при базе >0; число хранится долей, а не процентными пунктами. При базе 0 NULL.',
'known_ordered_units':'Число полученных уникальных событий-единиц. При неполной/устаревшей загрузке это известный подытог, не полный факт.',
'gross_drop':'SUM(MAX(base_14_avg − факт,0)) по сопоставимым карточкам; не доказанные потерянные заказы.',
'gross_gain':'SUM(MAX(факт − base_14_avg,0)) по сопоставимым карточкам; знаменатель долей роста отдельный от спада.',
'share_gross_drop':'Снижение карточки / валовой спад. Чистое изменение не используется как знаменатель.',
'scenario_units':'Отдельная условная арифметическая оценка по method/inputs. Не прогноз. Без подтверждённой сопоставимой доступности NULL.',
'orders_status':'OK по полному контракту загрузки; PARTIAL/ERROR/NO_ACCESS/MISSING/STALE/NOT_APPLICABLE различаются. OK не доказывает бизнес-причину.',
'snapshot_at':'Время исторического среза, если предоставлено источником; иначе время наблюдения current API. См. snapshot_time_basis. Всегда с явным часовым поясом.',
'source_run_id':'Ссылка на run_id в SourceStatus с теми же cabinet_id/source. Не идентификатор покупателя.',
'ordered_units':'Одна уникальная строка Statistics orders по srid = одна заказанная единица. Отменённые единицы не исключаются. Исходный srid не экспортируется.',
'cart_rate':'SUM(carts)/SUM(views) за один период и область. Не когортная вероятность и не среднее дневных процентов. Числитель и знаменатель сохранены.',
'order_rate':'SUM(funnel_orders)/SUM(carts) из Analytics одного периода/области. Знаменатель 0 → NULL. Не смешивать с заказами Statistics.',
'day':'Завершённый календарный день Europe/Moscow. Семантика события зависит от метрики: дата заказа или продажи указана в реестре.',
'observation_key':'Технический стабильный ключ объявленной области наблюдения. Не служит для соединения разных уровней детализации.',
'attribution_window':'Окно и метод атрибуции источника. Если точное окно WB не проверено, это явно подписано; органика вычитанием не вычисляется.',
'app_type':'Платформа из days.apps источника рекламы. Расход берётся только из листьев nms; итоги кампании не копируются на каждый nmId.',
'model':'Условные FBO/FBS по warehouseType/stockType WB: склад WB / склад продавца. Отдельная модель доставки не выдумывается; неизвестный warehouseType = UNKNOWN. Источник не позволяет отдельно исключить DBS без дополнительного подтверждения.',
'value':'Числовое значение с типом price_type/единицей из строки. Для снимков цены LAST; SUM по датам запрещён.',
}
EMPTY={
'Products':['cabinet_id','nm_id','supplier_article','title','life_status','first_observed_at'],
'ProductDaily':['cabinet_id','nm_id','day','ordered_units','orders_status','order_amount','amount_status','views','views_status','carts','carts_status','funnel_orders','wishlist'],
'VariantDaily':['cabinet_id','nm_id','barcode','size','model','day','ordered_units','status'],
'PriceHistory':['cabinet_id','nm_id','size_id','price_type','snapshot_at','value','currency','status'],
'StockSnapshots':['cabinet_id','nm_id','variant_id','warehouse_id','model','snapshot_at','quantity','status'],
'AdsDaily':['cabinet_id','campaign_id','nm_id','day','spend','impressions','clicks','attributed_orders','attributed_amount','attribution_window','status'],
'CampaignProducts':['cabinet_id','campaign_id','nm_id','valid_from','source','status'],
'CampaignSnapshots':['cabinet_id','campaign_id','snapshot_at','campaign_status','source_change_time','status'],
'SearchDaily':['cabinet_id','nm_id','period_start','period_end','query_scope','avg_position','status'],
'CardEvents':['cabinet_id','nm_id','observed_at','event_id','event_type','source','status'],
'CompetitorSnapshots':['competitor_nm_id','snapshot_at','source','price','currency','rating','status'],
'ProductMapping':['canonical_product_id','cabinet_id','nm_id','variant_id','units_per_item','comparison_unit','mapping_source','mapping_status'],
'Economics':['cabinet_id','nm_id','valid_from','unit_margin','source','status'],
'Evidence':['evidence_id','rule_id','rule_version','cabinet_id','nm_id','fact','next_action'],
'Actions':['action_id','evidence_id','status','owner','due_date','completed_at','metric_success'],
'PeriodMetrics':['cabinet_id','nm_id','window_days','period_start','period_end','previous_start','previous_end','units','previous_units','change_rate','status'],
'SourceStatus':['run_id','cabinet_id','source','status','reason','requested_at','completed_at','pages','records','cursor','snapshot_at','last_success_at'],
}


def headers(rows,name=''):
    fields=list(EMPTY.get(name,[]))
    for row in rows:
        for field in row:
            if field not in fields:fields.append(field)
    return fields or ['status','reason']


def value(v,field):
    if isinstance(v,(dict,list,tuple)):return json.dumps(v,ensure_ascii=False,default=str)
    if v is None:return None
    if field.endswith('_id') or field in {'nm_id','nm_a','nm_b','barcode','size_id','variant_id','warehouse_id','source_run_id'}:return str(v)
    if isinstance(v,str) and (field in {'day','report_date','period_start','period_end','previous_start','previous_end','valid_from','review_date','due_date'}):
        try:return date.fromisoformat(v[:10])
        except ValueError:pass
    return v


def table(wb,name,rows,fields=None,owner=False):
    ws=wb.create_sheet(name)
    fs=fields or headers(rows,name)
    ws.append([LABELS.get(f,f) if owner else f for f in fs])
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='233C56');cell.font=Font(name='Arial',size=10,color='FFFFFF',bold=True)
        cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    ws.row_dimensions[1].height=48 if owner else 30
    body_font=Font(name='Arial',size=10)
    for row_index,row in enumerate(rows,2):
        vals=[value(row.get(field),field) for field in fs]
        if owner:vals=['Н/Д' if v is None else v for v in vals]
        if owner and 'cabinet_id' in fs:
            from config import CABINETS
            idx=fs.index('cabinet_id');cab=CABINETS.get(str(vals[idx]));vals[idx]=cab.name if cab else vals[idx]
        ws.append(vals)
        for i,field in enumerate(fs,1):
            cell=ws.cell(row_index,i)
            if isinstance(cell.value,str):
                cell.data_type='s'  # Neutralize formula injection from vendor codes/imports.
                if owner and field in {'next_action','fact','scenario_limits','reason'}:cell.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(cell.value,(date,datetime)):cell.number_format='dd.mm.yyyy'
            elif isinstance(cell.value,(float,int)) and not isinstance(cell.value,bool):cell.number_format='0.0%' if field in PERCENT_FIELDS else '#,##0.00' if isinstance(cell.value,float) else '#,##0'
            cell.font=body_font
            if field=='wb_url' and cell.value:cell.hyperlink=cell.value;cell.style='Hyperlink'
        if owner and any(isinstance(row.get(f),str) and row.get(f) in {'ERROR','NO_ACCESS','PARTIAL','MISSING','STALE'} for f in fs):
            for i in range(1,len(fs)+1):ws.cell(row_index,i).fill=PatternFill('solid',fgColor='FFF3DF')
        if owner and any(f in fs for f in ('next_action','reason','fact')):ws.row_dimensions[row_index].height=60
    ws.freeze_panes='C2' if len(fs)>3 and 'nm_id' in fs else 'A2'
    ws.auto_filter.ref=ws.dimensions
    ws.sheet_view.showGridLines=False
    for j,f in enumerate(fs,1):
        width=50 if f in {'next_action','scenario_limits','reason','fact','limits','definition','missing_policy'} else 28 if f.endswith('_at') or f.endswith('_status') or f=='supplier_article' else 21 if owner else 20
        ws.column_dimensions[get_column_letter(j)].width=width
    if rows:
        tab=Table(displayName='T'+str(len(wb.worksheets)),ref=ws.dimensions)
        tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False)
        ws.add_table(tab)
    return ws


def dictionary(report):
    known={r['metric_id']:r for r in report['tables']['MetricRegistry']}
    result=[]
    for name,rows in report['tables'].items():
        for f in headers(rows,name):
            metric=known.get(f,{})
            result.append({'table_name':name,'field_id':f,'russian_name':LABELS.get(f,metric.get('russian_name',f)),
                           'definition':DEFINITIONS.get(f,metric.get('limits','Атрибут объявленного уровня строки; точный уникальный ключ в grain, источник/время в соседних полях и Manifest. Не самостоятельный бизнес-итог.')),
                           'unit':metric.get('unit','ID' if f.endswith('_id') else 'ISO 8601 с поясом' if f.endswith('_at') else 'date' if f in {'day','report_date'} else 'поле'),
                           'grain':TABLES.get(name,name),'source':metric.get('source',name),'aggregation':metric.get('aggregation','Не агрегировать без определения метрики'),
                           'time_semantics':metric.get('time_semantics','См. период/время и source_run_id строки')})
    return result


def overview(wb,report):
    ws=table(wb,'Обзор',report['summary'],['cabinet_id','ordered_units','known_ordered_units','base_14_avg','delta_units','change_rate','order_amount','gross_drop','gross_gain','available_days','expected_days','status'],True)
    ws.insert_rows(1,3)
    ws['A2']='ДЕМОНСТРАЦИЯ. ДАННЫЕ СИНТЕТИЧЕСКИЕ.' if report.get('data_mode') else 'Заказы Wildberries: завершённый день'
    ws['A2'].font=Font(name='Arial',size=14,bold=True,color='9C2525' if report.get('data_mode') else '233C56')
    ws.row_dimensions[1].height=8;ws.row_dimensions[2].height=24;ws.row_dimensions[3].height=8;ws.row_dimensions[4].height=48
    for t in ws.tables.values():t.ref=f'A4:L{len(report["summary"])+4}'
    ws.auto_filter.ref=f'A4:L{len(report["summary"])+4}';ws.freeze_panes='B5'
    # Context and limitations are visible directly above the charts, not hidden in metadata.
    row=ws.max_row+3
    notes=[f"Дата {report['report_date']}. Сформировано {report['generated_at']}. Версия {report['report_version']}, ревизия {report['revision']}.",
           f"База {report['intervals']['baseline_start']}–{report['intervals']['baseline_end']}. Статус пакета: {report['status']}.",
           'Н/Д в обзорах и пустые числовые ячейки в исходных таблицах означают отсутствие данных. Известный подытог не является полным итогом. Заказы — заказанные единицы, отмены включены.',
           'FBO: warehouseType=Склад WB; FBS: Склад продавца. Неизвестная модель хранится отдельно.',
           *report['limitations']]
    for text in notes:
        ws.cell(row,1,text);ws.cell(row,1).alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row].height=max(30,((len(text)+59)//60)*16);row+=1
    # Wide single note column uses no merge; chart starts below the context.
    ws.column_dimensions['A'].width=66
    trend=defaultdict(list)
    for r in report['tables']['ProductDaily']:trend[str(r['day'])[:10]].append(r)
    data=[]
    for day,rows in sorted(trend.items()):
        valid=all(r['orders_status'] in {'OK','NOT_APPLICABLE'} and r.get('source_day_covered',False) for r in rows)
        nums=[r['ordered_units'] for r in rows if r['orders_status']=='OK']
        money=[r['order_amount'] for r in rows if r['orders_status']=='OK']
        units=sum(nums) if valid else None
        amount=sum(money) if valid and all(v is not None for v in money) else None
        data.append({'day':day,'ordered_units':units,'avg_order_price':amount/units if amount is not None and units else None,'status':'OK' if valid else 'PARTIAL'})
    history=table(wb,'Дневной график',data,['day','ordered_units','avg_order_price','status'])
    if data:
        chart=LineChart();chart.title='Заказы и средняя цена заказов';chart.y_axis.title='Заказано, шт.';chart.x_axis.title='Завершённый день'
        chart.x_axis=DateAxis(crossAx=100);chart.x_axis.axId=10;chart.x_axis.number_format='dd.mm';chart.x_axis.majorTimeUnit='days';chart.x_axis.title='Завершённый день'
        chart.add_data(Reference(history,min_col=2,min_row=1,max_row=history.max_row),titles_from_data=True)
        chart.set_categories(Reference(history,min_col=1,min_row=2,max_row=history.max_row));chart.display_blanks='gap'
        price=LineChart();price.add_data(Reference(history,min_col=3,min_row=1,max_row=history.max_row),titles_from_data=True)
        price.set_categories(Reference(history,min_col=1,min_row=2,max_row=history.max_row))
        price.y_axis.axId=200;price.y_axis.title='Средняя цена заказа, ₽';price.y_axis.crosses='max';price.display_blanks='gap'
        chart+=price;chart.width=26;chart.height=11;ws.add_chart(chart,f'A{row+2}')


def workbook(report,full=True):
    wb=Workbook();wb.remove(wb.active)
    overview(wb,report)
    fields=['cabinet_id','nm_id','supplier_article','ordered_units','base_14_avg','delta_units','change_rate','signal','orders_last_7_sum','week_change_rate','weekday_avg','weekday_available','order_amount','orders_status','fbo_snapshot_quantity','fbo_status','fbs_snapshot_quantity','fbs_status','stock_coverage_days_estimate','wb_url']
    table(wb,'Все товары',report['all_products'],fields,True)
    drops=sorted([r for r in report['all_products'] if r['legacy_drop_selected']],key=lambda r:r['priority_absolute'] or 0,reverse=True)
    table(wb,'Падения и действия',[{**r,'next_action':next((e['next_action'] for e in report['tables']['Evidence'] if e['cabinet_id']==r['cabinet_id'] and e['nm_id']==r['nm_id']),'Причина не установлена; проверить данные')} for r in drops],fields[:8]+['share_gross_drop','alert_selected','next_action'],True)
    growth=[e for e in report['tables']['Evidence'] if e['category']=='Проверить возможность роста']
    table(wb,'Рост и действия',growth,['cabinet_id','nm_id','fact','explanation_status','next_action','scenario_kind','scenario_units','scenario_horizon_days','scenario_limits'],True)
    comparison=table(wb,'Сравнение кабинетов',report['comparisons'],['canonical_product_id','comparison_unit','fact_units','base_units','delta_units','change_rate','cabinet_a','nm_a','cabinet_b','nm_b','status','limit'],True)
    if not report['comparisons']:
        comparison['A3']='Нет подтверждённых сопоставлений. Импортируйте ProductMapping с единицами сравнения и комплектностью.'
        comparison['A3'].alignment=Alignment(wrap_text=True,vertical='top');comparison.row_dimensions[3].height=100;comparison.column_dimensions['A'].width=50
    quality=list(report['tables']['SourceStatus'])
    for s in report['summary']:
        quality.append({'cabinet_id':s['cabinet_id'],'source':'coverage_orders','status':s['status'],'reason':s['reason'],'records':s['products'],'completed_at':report['generated_at']})
    table(wb,'Качество данных',quality,['cabinet_id','source','status','reason','records','requested_at','completed_at','last_success_at','pages','cursor'],True)
    manifest=[{'field_id':k,'value':v} for k,v in report.items() if k not in {'tables','all_products','summary','comparisons'}]
    for name,rows in report['tables'].items():manifest.append({'field_id':'table:'+name,'value':{'rows':len(rows),'grain':TABLES.get(name),'status':'AVAILABLE' if rows else 'MISSING','location':name if full else name+'.csv'}})
    table(wb,'Manifest',manifest)
    table(wb,'Dictionary',dictionary(report))
    if full:
        for name,rows in report['tables'].items():table(wb,name,rows)
    stream=BytesIO();wb.save(stream);return stream.getvalue()


def build_package(report,max_bytes=None,max_cells=None):
    max_bytes=max_bytes if max_bytes is not None else int(os.getenv('REPORT_MAX_BYTES',str(45*1024*1024)))
    max_cells=max_cells if max_cells is not None else int(os.getenv('REPORT_MAX_CELLS','450000'))
    cells=sum((len(rows)+1)*len(headers(rows,name)) for name,rows in report['tables'].items())
    too_large=cells>max_cells or any(len(rows)>1048575 for rows in report['tables'].values())
    date_str=report['report_date'];version=report['report_version'];revision=report['revision']
    prefix=f'WB_аналитика_{date_str}_v{version}_r{revision}'
    if not too_large:
        xlsx=workbook(report)
        if len(xlsx)<=max_bytes:return prefix+'.xlsx',xlsx
    manifest={k:v for k,v in report.items() if k not in {'tables','all_products','summary','comparisons'}}
    manifest.update(package_reason='Превышен предел строк/ячеек или размер XLSX; полный ассортимент и история сохранены в CSV',csv_encoding='utf-8-sig',csv_delimiter=';',decimal_separator='.',tables={})
    out=BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('Обзор.xlsx',workbook(report,full=False))
        tables={**report['tables'],'Dictionary':dictionary(report),'AllProducts':report['all_products'],'Summary':report['summary'],'CabinetComparison':report['comparisons']}
        for name,rows in tables.items():
            stream=StringIO(newline='');fs=headers(rows,name)
            writer=csv.DictWriter(stream,fieldnames=fs,delimiter=';',lineterminator='\n');writer.writeheader()
            for row in rows:
                record={f:value(row.get(f),f) for f in fs}
                # CSV is external machine data: quote leading formula characters for spreadsheet safety.
                for f,v in record.items():
                    if isinstance(v,str) and v.startswith(('=','+','-','@')):record[f]="'"+v
                writer.writerow(record)
            raw=stream.getvalue().encode('utf-8-sig');z.writestr(name+'.csv',raw)
            manifest['tables'][name]={'rows':len(rows),'grain':TABLES.get(name),'sha256':hashlib.sha256(raw).hexdigest(),'fields':fs}
        z.writestr('manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2,default=str))
        z.writestr('README.txt','Пересылайте весь ZIP. Обзор.xlsx содержит результаты. CSV — полные таблицы без обрезки. NULL — пустое поле + статус. Разделитель ;, десятичный знак ., UTF-8 BOM. Строки с ведущими =+-@ имеют защитный апостроф.\n')
    payload=out.getvalue()
    if len(payload)>max_bytes:raise ValueError('Полный ZIP превышает разрешённый размер отправки; данные сохранены в БД. Увеличьте разрешённый размер/скачайте пакет на сервере. История не обрезана.')
    return prefix+'.zip',payload
