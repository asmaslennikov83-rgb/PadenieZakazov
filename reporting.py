from __future__ import annotations

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from analysis_service import CabinetReport


def format_number(value: float) -> str:
    if abs(value - round(value)) < 0.005:
        return str(int(round(value)))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def format_report_message(report: CabinetReport, limit: int = 20) -> str:
    if not report.rows:
        return (
            f"✅ <b>{report.cabinet_name}</b>\n"
            f"За {report.yesterday:%d.%m.%Y} падения заказов "
            f"более {report.threshold}% не обнаружено."
        )

    lines = [
        "📉 <b>ПАДЕНИЕ ЗАКАЗОВ WILDBERRIES</b>",
        "",
        f"📅 Вчера: <b>{report.yesterday:%d.%m.%Y}</b>",
        f"🏢 Кабинет: <b>{report.cabinet_name}</b>",
        (
            f"📊 Среднее: {report.baseline_start:%d.%m.%Y}–"
            f"{report.baseline_end:%d.%m.%Y}"
        ),
        f"⚙️ Порог: более {report.threshold}%",
        "",
    ]

    for index, row in enumerate(report.rows[:limit], start=1):
        icon = "🔴" if row.yesterday_orders == 0 else "🔻"
        lines.extend(
            [
                f"<b>{index}. {html_escape(row.supplier_article)}</b>",
                f"Среднее в день: {format_number(row.average_per_day)} шт.",
                f"Заказов вчера: {row.yesterday_orders} шт.",
                f"{icon} Падение: {format_number(row.drop_percent)}%",
                "",
            ]
        )

    if len(report.rows) > limit:
        lines.append(f"Показано: {limit} из {len(report.rows)} позиций.")
    lines.append(f"Всего артикулов с падением: <b>{len(report.rows)}</b>")

    return "\n".join(lines)


def html_escape(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def build_excel(reports: list[CabinetReport]) -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    headers = [
        "Кабинет",
        "Артикул продавца",
        "Среднее заказов в день",
        "Заказов вчера",
        "Падение, %",
        "Абсолютное снижение",
        "Базовый период с",
        "Базовый период по",
        "Дата сравнения",
    ]

    for report in reports:
        title = sanitize_sheet_name(report.cabinet_name)
        ws = wb.create_sheet(title=title)

        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for row in report.rows:
            ws.append(
                [
                    report.cabinet_name,
                    row.supplier_article,
                    round(row.average_per_day, 2),
                    row.yesterday_orders,
                    round(row.drop_percent, 2),
                    round(row.absolute_drop, 2),
                    report.baseline_start.strftime("%d.%m.%Y"),
                    report.baseline_end.strftime("%d.%m.%Y"),
                    report.yesterday.strftime("%d.%m.%Y"),
                ]
            )

        widths = [24, 28, 24, 18, 16, 22, 18, 18, 18]
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    if not wb.sheetnames:
        ws = wb.create_sheet("Отчёт")
        ws["A1"] = "Нет данных"

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def sanitize_sheet_name(value: str) -> str:
    bad = '[]:*?/\\'
    for ch in bad:
        value = value.replace(ch, " ")
    value = value.strip() or "Отчёт"
    return value[:31]
